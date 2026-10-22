from datetime import datetime, timezone
from pathlib import Path
from secrets import token_urlsafe
from time import time
import os
import re
import threading

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fastapi import FastAPI, HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field

app = FastAPI(title="Resume Tailor DOCX Generator")

OUTPUT_DIR = Path("/tmp/resume-tailor-docx")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
DOWNLOADS = {}
LOCK = threading.Lock()
RETENTION_SECONDS = 2 * 60 * 60
PUBLIC_BASE_URL = os.getenv("DOCX_PUBLIC_BASE_URL", "http://localhost:8000").rstrip("/")

NAVY = (31, 58, 95)
RULE_BLUE = (65, 92, 135)
BODY_COLOR = (35, 35, 35)


class GenerateRequest(BaseModel):
    tailored_resume: str = Field(min_length=1)
    review_report: dict = Field(default_factory=dict)


def clean_text(value):
    return "" if value is None else str(value).strip()


def set_run_font(run, size=10, bold=False, italic=False, color=BODY_COLOR):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color)


def add_bottom_border(paragraph, color="415C87", size="8", space="4"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_borders = p_pr.find(qn("w:pBdr"))
    if p_borders is None:
        p_borders = OxmlElement("w:pBdr")
        p_pr.append(p_borders)

    bottom = p_borders.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_borders.append(bottom)

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)


def add_text_paragraph(doc, text="", size=9.5, bold=False, italic=False,
                       before=0, after=2, align=None, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = keep_with_next
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def looks_like_contact(line):
    lower = line.lower()
    return (
        "@" in line
        or re.search(r"\b(?:\+?\d[\d\s().-]{7,}\d)\b", line) is not None
        or "linkedin.com" in lower
        or "github.com" in lower
        or "portfolio" in lower
        or "https://" in lower
        or "http://" in lower
    )


HEADING_ALIASES = {
    "summary": "SUMMARY",
    "professional summary": "SUMMARY",
    "profile": "SUMMARY",
    "objective": "SUMMARY",
    "work experience": "PROFESSIONAL EXPERIENCE",
    "professional experience": "PROFESSIONAL EXPERIENCE",
    "experience": "PROFESSIONAL EXPERIENCE",
    "employment history": "PROFESSIONAL EXPERIENCE",
    "internship experience": "PROFESSIONAL EXPERIENCE",
    "internships": "PROFESSIONAL EXPERIENCE",
    "education": "EDUCATION",
    "academic background": "EDUCATION",
    "projects": "PROJECTS",
    "academic projects": "PROJECTS",
    "selected projects": "PROJECTS",
    "technical skills": "SKILLS",
    "skills": "SKILLS",
    "core skills": "SKILLS",
    "skills & languages": "SKILLS & LANGUAGES",
    "skills and languages": "SKILLS & LANGUAGES",
    "certifications": "CERTIFICATIONS",
    "certificates": "CERTIFICATIONS",
    "licenses": "CERTIFICATIONS",
    "awards": "AWARDS",
    "languages": "LANGUAGES",
    "volunteer experience": "VOLUNTEER EXPERIENCE",
    "activities": "ACTIVITIES",
    "publications": "PUBLICATIONS",
    "references": "REFERENCES",
}




def heading_label(line):
    normalized = re.sub(r"[:\s]+$", "", line.strip()).lower()

    # Only recognized section headings are styled as headings.
    # This prevents all-caps skill names such as HTML, CSS, SQL,
    # or PHP from becoming accidental section headings.
    return HEADING_ALIASES.get(normalized)



def is_bullet(line):
    return bool(re.match(r"^\s*(?:[-*•▪◦‣]|\d+[.)])\s+", line))


def strip_bullet(line):
    return re.sub(r"^\s*(?:[-*•▪◦‣]|\d+[.)])\s+", "", line).strip()


def add_section_heading(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    run = p.add_run(label)
    set_run_font(run, size=11, bold=True, color=NAVY)
    add_bottom_border(p, color="415C87", size="8", space="4")


def _split_labeled_line(line):
    # Bold a leading label before a colon (e.g. "CGPA: 3.42 / 4.00").
    match = re.match(r"^([^:]{2,38}:)(\s*)(.*)$", line)
    if match:
        return match.group(1), match.group(3)
    return None


def add_bullet(doc, line, size=9.5):
    text = strip_bullet(line)
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_together = True

    labeled = _split_labeled_line(text)
    if labeled:
        lead, rest = labeled
        r1 = p.add_run(lead)
        set_run_font(r1, size=size, bold=True)
        if rest:
            r2 = p.add_run(" " + rest)
            set_run_font(r2, size=size)
    else:
        run = p.add_run(text)
        set_run_font(run, size=size)
    return p


def format_resume(doc, text):
    lines = [line.strip() for line in clean_text(text).replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    lines = [line for line in lines if line]
    if not lines:
        raise HTTPException(status_code=422, detail="The tailored resume was empty.")

    # Name and contact details: centered and visually distinct.
    add_text_paragraph(
        doc, lines[0], size=19, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
        before=0, after=4, keep_with_next=True
    )

    start_index = 1
    contact_index = None
    for i in range(1, min(4, len(lines))):
        if looks_like_contact(lines[i]):
            contact_index = i
            break

    if contact_index is not None:
        add_text_paragraph(
            doc, lines[contact_index], size=9, align=WD_ALIGN_PARAGRAPH.CENTER,
            after=9, keep_with_next=True
        )
        lines.pop(contact_index)
    else:
        # Leave a modest gap after the name when no contact line is supplied.
        doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

    for line in lines[start_index:]:
        label = heading_label(line)
        if label:
            add_section_heading(doc, label)
            continue

        if is_bullet(line):
            add_bullet(doc, line)
            continue

        # Format role/project/school lines with a stronger visual hierarchy.
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.keep_together = True

        likely_title = (
            len(line) <= 115
            and len(line.split()) <= 18
            and not line.endswith((".", ";", ","))
            and any(ch.isalpha() for ch in line)
            and not looks_like_contact(line)
        )
        run = p.add_run(line)
        set_run_font(run, size=9.5, bold=likely_title)


def report_items(report, key):
    value = report.get(key, [])
    return value if isinstance(value, list) else []


def add_report(doc, report):
    doc.add_page_break()
    add_text_paragraph(doc, "REVIEW REPORT", size=16, bold=True, after=8, keep_with_next=True)

    sections = [
        ("Changes Made", "changes_made"),
        ("Skills or Experience to Confirm", "skills_to_confirm"),
        ("Other Recommendations", "other_recommendations"),
    ]

    for heading, key in sections:
        add_section_heading(doc, heading.upper())
        items = report_items(report, key)
        if not items:
            add_text_paragraph(doc, "None listed.", size=9.5, after=3)
            continue

        for item in items:
            if isinstance(item, dict):
                skill = clean_text(item.get("skill"))
                why = clean_text(item.get("why_it_matters"))
                question = clean_text(item.get("question"))
                parts = []
                if skill:
                    parts.append(skill)
                if why:
                    parts.append(f"Why it matters: {why}")
                if question:
                    parts.append(question)
                text = " — ".join(parts) if parts else str(item)
            else:
                text = clean_text(item)
            add_bullet(doc, text, size=9.5)


def remove_expired():
    now = time()
    with LOCK:
        expired = [token for token, entry in DOWNLOADS.items() if entry["expires_at"] <= now]
        for token in expired:
            entry = DOWNLOADS.pop(token)
            try:
                Path(entry["path"]).unlink(missing_ok=True)
            except OSError:
                pass


@app.get("/health")
def health():
    return {"status": "ok"}


@app.post("/generate")
def generate(payload: GenerateRequest):
    remove_expired()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(*BODY_COLOR)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0

    format_resume(doc, payload.tailored_resume)
    if payload.review_report:
        add_report(doc, payload.review_report)

    token = token_urlsafe(24)
    filename = f"tailored_resume_{datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')}.docx"
    filepath = OUTPUT_DIR / f"{token}.docx"
    doc.save(filepath)

    with LOCK:
        DOWNLOADS[token] = {
            "path": str(filepath),
            "filename": filename,
            "expires_at": time() + RETENTION_SECONDS,
        }

    return {
        "filename": filename,
        "download_url": f"{PUBLIC_BASE_URL}/download/{token}",
        "review_report": payload.review_report,
        "expires_in_seconds": RETENTION_SECONDS,
    }


@app.get("/download/{token}")
def download(token: str):
    remove_expired()
    with LOCK:
        entry = DOWNLOADS.get(token)

    if not entry or entry["expires_at"] <= time():
        raise HTTPException(
            status_code=404,
            detail="This download link has expired or is invalid. Please submit the form again.",
        )

    path = Path(entry["path"])
    if not path.is_file():
        raise HTTPException(
            status_code=404,
            detail="The requested document is no longer available. Please submit the form again.",
        )

    return FileResponse(
        path,
        filename=entry["filename"],
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
